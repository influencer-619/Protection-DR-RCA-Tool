"""PDF / Word text extract → settings ingest."""

from __future__ import annotations

import io

from app.services.document_extract import extract_document_text, extract_docx_text, extract_pdf_text
from app.services.file_service import infer_source_type
from app.services.settings_ingest import ingest_settings_bytes


SETTINGS_BODY = """Relay settings export
Phase CT Primary    2500
Phase CT Secondary    1
Main VT Primary    6600
Main VT Secondary    110
I>1 Current Set    1.8
I>1 TMS    0.16
pickup_current=1.8
ct_ratio=2500
vt_ratio=60
time_dial=0.16
curve=IEC S Inverse
"""


def _make_pdf(text: str) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    y = 800
    for line in text.splitlines():
        c.drawString(40, y, line[:110])
        y -= 14
        if y < 40:
            c.showPage()
            y = 800
    c.save()
    return buf.getvalue()


def _make_docx(text: str) -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading("Protection settings", level=1)
    for line in text.splitlines():
        if "=" in line:
            k, _, v = line.partition("=")
            table = doc.add_table(rows=1, cols=2)
            table.rows[0].cells[0].text = k.strip()
            table.rows[0].cells[1].text = v.strip()
        else:
            doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_pdf_settings_extract_and_ingest():
    pdf = _make_pdf(SETTINGS_BODY)
    extracted = extract_pdf_text(pdf)
    assert extracted["status"] == "OK"
    assert "pickup_current" in extracted["text"] or "1.8" in extracted["text"]

    ing = ingest_settings_bytes(pdf, filename="feeder_settings.pdf")
    assert ing["status"] == "OK"
    assert ing.get("param_count", 0) > 0
    common = ing.get("common") or {}
    mapped = ing.get("mapped") or {}
    assert common.get("pickup_current") == 1.8 or (mapped.get("ct_vt") or {}).get("ct_ratio") or common.get(
        "ct_ratio"
    )


def test_docx_settings_extract_and_ingest():
    docx = _make_docx(SETTINGS_BODY)
    extracted = extract_docx_text(docx)
    assert extracted["status"] == "OK"
    assert "ct_ratio" in extracted["text"].lower() or "2500" in extracted["text"]

    ing = ingest_settings_bytes(docx, filename="transformer_settings.docx")
    assert ing["status"] == "OK"
    assert (ing.get("raw") or {}).get("document_format") == "docx"


def test_infer_source_type_document_settings_and_events():
    assert infer_source_type("P143_relay_settings.pdf") == "SETTINGS"
    assert infer_source_type("unit_settings.docx") == "SETTINGS"
    assert infer_source_type("fault_event_report.pdf") == "RELAY_EVENT_REPORT"
    assert infer_source_type("notes_only.pdf") == "ATTACHMENT"


def test_empty_pdf_not_calculable():
    # Minimal PDF with almost no text
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.showPage()
    c.save()
    empty = buf.getvalue()
    extracted = extract_document_text(empty, filename="blank.pdf")
    assert extracted["status"] == "NOT_CALCULABLE"
    ing = ingest_settings_bytes(empty, filename="blank_settings.pdf")
    assert ing["status"] == "NOT_CALCULABLE"
