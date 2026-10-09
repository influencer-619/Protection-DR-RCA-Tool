"""One-shot: export Complete RCA Matrix Excel → JSON for matrix_v1 generation."""

from __future__ import annotations

import json
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("openpyxl required", file=sys.stderr)
    raise

SRC = Path(
    r"C:\Users\5863.AVAADA\Downloads\Complete_Generalized_Protection_RCA_Matrix_L1_L2_L3_Fallback_LBB.xlsx"
)
OUT = Path(__file__).resolve().parents[2] / "rules" / "rca" / "_excel_matrix_export.json"


def main() -> None:
    wb = openpyxl.load_workbook(SRC, data_only=True)
    ws = wb["Complete RCA Matrix"]
    rows = []
    for r in range(2, (ws.max_row or 0) + 1):
        scenario = ws.cell(r, 2).value
        if not scenario:
            continue
        rows.append(
            {
                "row": r,
                "asset": str(ws.cell(r, 1).value or "").strip(),
                "scenario": str(scenario).strip(),
                "ansi": str(ws.cell(r, 3).value or "").strip(),
                "l1": str(ws.cell(r, 4).value or "").strip(),
                "l2": str(ws.cell(r, 5).value or "").strip(),
                "l3": str(ws.cell(r, 6).value or "").strip(),
                "fallback": str(ws.cell(r, 7).value or "").strip(),
                "final_class": str(ws.cell(r, 8).value or "").strip(),
            }
        )
    lbb = []
    ws2 = wb["LBB Cascade Logic"]
    for r in range(2, (ws2.max_row or 0) + 1):
        lbb.append(
            {
                "step": ws2.cell(r, 1).value,
                "question": ws2.cell(r, 2).value,
                "if_yes": ws2.cell(r, 3).value,
                "if_no": ws2.cell(r, 4).value,
                "meaning": ws2.cell(r, 5).value,
            }
        )
    fallback = []
    ws3 = wb["Fallback Methodology"]
    for r in range(2, (ws3.max_row or 0) + 1):
        fallback.append(
            {
                "level": ws3.cell(r, 1).value,
                "primary": ws3.cell(r, 2).value,
                "fallback": ws3.cell(r, 3).value,
                "rule": ws3.cell(r, 4).value,
            }
        )
    taxonomy = []
    ws4 = wb["RCA Taxonomy"]
    for r in range(2, (ws4.max_row or 0) + 1):
        taxonomy.append(
            {
                "classification": ws4.cell(r, 1).value,
                "meaning": ws4.cell(r, 2).value,
                "minimum": ws4.cell(r, 3).value,
            }
        )
    payload = {
        "source": SRC.name,
        "scenarios": rows,
        "lbb_cascade_logic": lbb,
        "fallback_methodology": fallback,
        "rca_taxonomy": taxonomy,
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(payload, indent=2), encoding="utf-8")
    print(f"scenarios={len(rows)} wrote {OUT}")
    for i, x in enumerate(rows, 1):
        print(f"{i:02d}|{x['asset']}|{x['scenario']}|{x['final_class'][:70]}")


if __name__ == "__main__":
    main()
